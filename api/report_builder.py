"""
report_builder.py — generate PDF and DOCX investigation reports.
PDF:   reportlab (pure Python)
DOCX:  python-docx (pure Python)
"""
from __future__ import annotations
import json, os, re, tempfile
from datetime import datetime
from typing import TYPE_CHECKING
if TYPE_CHECKING:
    from api.models import Case

_CYAN   = (0x39, 0xD0, 0xD8)
_ORANGE = (0xF0, 0x88, 0x3E)
_GREEN  = (0x3F, 0xB9, 0x50)
_RED    = (0xF8, 0x51, 0x49)
_PURPLE = (0xBC, 0x8C, 0xFF)
_BLUE   = (0x21, 0x96, 0xF3)
_YELLOW = (0xFF, 0xD5, 0x4F)
_LIGHT  = (0x00, 0x00, 0x00)
_MUTED  = (0x00, 0x00, 0x00)

def _ts(dt):
    return dt.strftime("%Y-%m-%d %H:%M UTC") if dt else "—"

def _sev(s):
    return (s or "unknown").upper()


def _pipe_display(value: str | None) -> str:
    if not value:
        return "—"
    return ", ".join([part.strip() for part in value.split("|") if part.strip()]) or "—"


def _norm_tcode(value: str | None) -> str:
    text = str(value or "").strip().upper()
    match = re.search(r"T\d{4}(?:\.\d{3})?", text)
    return match.group(0) if match else text


def _bias_result(case: "Case") -> dict | None:
    if not case.bias_result_json:
        return None
    try:
        return json.loads(case.bias_result_json)
    except Exception:
        return None


def _connection_class(item: dict) -> str:
    score = float(item.get("transition_score", 0.0) or 0.0)
    conn_class = str(item.get("connection_class") or "").strip().lower()
    if conn_class:
        return conn_class
    if bool(item.get("is_gap", False)):
        return "strong_gap" if score and score < 0.35 else "weak_gap"
    if score >= 0.80:
        return "direct_connection"
    if score >= 0.65:
        return "weak_direct"
    return "indeterminate"


def _connection_label(item: dict) -> str:
    label = str(item.get("connection_label") or "").strip()
    if label:
        return label
    mapping = {
        "direct_connection": "Direct Connection",
        "weak_direct": "Weak Direct",
        "indeterminate": "Indeterminate",
        "weak_gap": "Weak Gap",
        "strong_gap": "Strong Gap",
    }
    return mapping.get(_connection_class(item), "Indeterminate")


def _case_meta_rows(case: "Case") -> list[list[str]]:
    return [
        ["Case Number", case.case_number or "—"],
        ["Title", case.title or "—"],
        ["Severity", _sev(case.severity.value if case.severity else "")],
        ["Status", (case.status.value if case.status else "").replace("_", " ").title()],
        ["Owner", case.owner.username if case.owner else "Unassigned"],
        ["Template ID", str(case.template_id) if case.template_id else "—"],
        ["Created", _ts(case.created_at)],
        ["Updated", _ts(case.updated_at)],
        ["BIAS Run", _ts(case.bias_ran_at)],
        ["BIAS Platforms", _pipe_display(case.bias_platforms)],
        ["BIAS CLI Environment", _pipe_display(case.bias_cli_env)],
        ["BIAS Scripting", _pipe_display(case.bias_scripting)],
        ["BIAS Network", _pipe_display(case.bias_network)],
    ]


def _entity_related_cases(case: "Case", case_entity) -> str:
    entity = getattr(case_entity, "entity", None)
    if not entity:
        return "—"
    current_case_id = getattr(case, "id", None)
    related = []
    seen = set()
    for link in getattr(entity, "case_links", []) or []:
        linked_case = getattr(link, "case", None)
        if not linked_case or getattr(linked_case, "id", None) == current_case_id:
            continue
        case_number = getattr(linked_case, "case_number", None) or "CASE"
        title = getattr(linked_case, "title", None) or "Untitled"
        label = f"{case_number} ({title})"
        if label in seen:
            continue
        seen.add(label)
        related.append(label)
    return ", ".join(related) if related else "None"


def _render_bias_graph(case: "Case", output_path: str) -> str | None:
    result = _bias_result(case)
    observables = list(case.observables or [])
    if not result or not observables:
        return None

    norm_obs = []
    seen = set()
    for obs in observables:
        tcode = _norm_tcode(obs.tcode)
        if not tcode or tcode in seen:
            continue
        seen.add(tcode)
        norm_obs.append({
            "tcode": tcode,
            "tname": obs.tname or "",
        })
    if not norm_obs:
        return None

    os.environ.setdefault("MPLCONFIGDIR", tempfile.gettempdir())
    import matplotlib
    matplotlib.use("Agg", force=True)
    import matplotlib.pyplot as plt
    from matplotlib.lines import Line2D

    x_step = 4.0
    top_base_y = 1.4
    bridge_gap = 0.9
    bottom_base_y = -1.8
    obs_x = {obs["tcode"]: i * x_step for i, obs in enumerate(norm_obs)}
    max_bridges = 0
    has_chain = False

    for gap in result.get("gaps", []):
        cands = gap.get("candidates", []) if isinstance(gap.get("candidates", []), list) else []
        chains = gap.get("chains", []) if isinstance(gap.get("chains", []), list) else []
        max_bridges = max(max_bridges, min(3, len(cands)))
        has_chain = has_chain or bool(chains[:1])

    fig_h = max(5.0, 2.8 + max_bridges * 0.85 + (1.5 if has_chain else 0))
    fig_w = max(10.0, len(norm_obs) * 3.2)
    fig, ax = plt.subplots(figsize=(fig_w, fig_h))
    fig.patch.set_facecolor("#FFFFFF")
    ax.set_facecolor("#FFFFFF")
    ax.axis("off")

    anchor_ids = set()
    gap_num = 0
    for i, gap in enumerate(result.get("gaps", [])):
        src = _norm_tcode(gap.get("from_tcode"))
        tgt = _norm_tcode(gap.get("to_tcode"))
        if src not in obs_x or tgt not in obs_x:
            continue
        anchor_ids.update([src, tgt])
        x_src, x_tgt = obs_x[src], obs_x[tgt]
        x_mid = (x_src + x_tgt) / 2
        candidates = (gap.get("candidates", []) if isinstance(gap.get("candidates", []), list) else [])[:3]
        chains = (gap.get("chains", []) if isinstance(gap.get("chains", []), list) else [])[:1]
        is_gap = bool(gap.get("is_gap", False))
        conn_class = _connection_class(gap)
        conn_label = _connection_label(gap)
        if is_gap:
            gap_num += 1
        style_map = {
            "direct_connection": ("#2196F3", "-", 2.6),
            "weak_direct": ("#2196F3", ":", 2.2),
            "indeterminate": ("#FFD54F", ":", 2.2),
            "weak_gap": ("#F0883E", ":", 2.2),
            "strong_gap": ("#F85149", "-", 2.8),
        }
        edge_color, edge_style, edge_width = style_map.get(conn_class, style_map["indeterminate"])
        edge_label = f"Gap {gap_num}" if is_gap else conn_label
        ax.plot([x_src, x_tgt], [0, 0], color=edge_color, linestyle=edge_style, linewidth=edge_width, zorder=1)
        ax.text(x_mid, 0.2, edge_label, color="#000000",
                ha="center", va="bottom", fontsize=9, fontweight="bold")

        if not is_gap:
            continue

        for r, cand in enumerate(candidates):
            y = top_base_y + (r * bridge_gap)
            cand_tcode = _norm_tcode(cand.get("tcode"))
            score = float(cand.get("score", 0) or 0)
            edge_col = "#3FB950" if score >= 0.8 else "#F0883E" if score >= 0.5 else "#F85149"
            ax.plot([x_src, x_mid], [0, y], color=edge_col, linewidth=1.8, zorder=1)
            ax.plot([x_mid, x_tgt], [y, 0], color=edge_col, linewidth=1.8, zorder=1)
            ax.scatter([x_mid], [y], s=420, marker="D", color="#bf360c", edgecolors="#F0883E", linewidths=1.2, zorder=3)
            ax.text(x_mid, y, f"{cand_tcode}\n{int(score * 100)}%", color="#000000", ha="center", va="center", fontsize=8, zorder=4)

        if chains:
            chain = chains[0]
            path = chain.get("path", []) if isinstance(chain.get("path", []), list) else []
            if path:
                prev_x, prev_y = x_src, 0
                for h_idx, hop in enumerate(path):
                    hop = _norm_tcode(hop)
                    hop_x = x_src + ((x_tgt - x_src) * ((h_idx + 1) / (len(path) + 1)))
                    hop_y = bottom_base_y
                    ax.plot([prev_x, hop_x], [prev_y, hop_y], color="#BC8CFF", linestyle="--", linewidth=1.8, zorder=1)
                    ax.scatter([hop_x], [hop_y], s=500, marker="h", color="#311b92", edgecolors="#BC8CFF", linewidths=1.5, zorder=3)
                    ax.text(hop_x, hop_y, hop, color="#000000", ha="center", va="center", fontsize=8, zorder=4)
                    prev_x, prev_y = hop_x, hop_y
                ax.plot([prev_x, x_tgt], [prev_y, 0], color="#BC8CFF", linestyle="--", linewidth=1.8, zorder=1)

    for obs in norm_obs:
        x = obs_x[obs["tcode"]]
        border = "#F0883E" if obs["tcode"] in anchor_ids else "#39D0D8"
        ax.scatter([x], [0], s=900, marker="s", color="#0097a7", edgecolors=border, linewidths=2.0, zorder=3)
        label = obs["tcode"] if not obs["tname"] else f"{obs['tcode']}\n{obs['tname']}"
        ax.text(x, 0, label, color="#000000", ha="center", va="center", fontsize=8, zorder=4)

    legend_items = [
        Line2D([0], [0], color="#2196F3", lw=2.2, linestyle="-", label="Direct / weak direct"),
        Line2D([0], [0], color="#F0883E", lw=2.2, linestyle=":", label="Weak gap"),
        Line2D([0], [0], color="#F85149", lw=2.6, linestyle="-", label="Strong gap"),
        Line2D([0], [0], marker="s", color="w", markerfacecolor="#0097a7", markeredgecolor="#39D0D8", markersize=10, label="Observable"),
        Line2D([0], [0], marker="D", color="w", markerfacecolor="#bf360c", markeredgecolor="#F0883E", markersize=9, label="Bridge candidate"),
        Line2D([0], [0], marker="h", color="#BC8CFF", markerfacecolor="#311b92", markeredgecolor="#BC8CFF", linestyle="--", markersize=10, label="Multi-hop path"),
    ]
    ax.legend(
        handles=legend_items,
        loc="upper center",
        bbox_to_anchor=(0.5, 1.06),
        ncol=3,
        frameon=False,
        fontsize=8,
        labelcolor="#000000",
    )

    xs = list(obs_x.values())
    ax.set_xlim(min(xs) - 1.8, max(xs) + 1.8)
    ax.set_ylim(bottom_base_y - 1.2, top_base_y + max_bridges * bridge_gap + 1.2)
    fig.tight_layout()
    fig.savefig(output_path, dpi=180, facecolor=fig.get_facecolor(), bbox_inches="tight")
    plt.close(fig)
    return output_path

def build_pdf(case: "Case", output_path: str) -> str:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                     TableStyle, HRFlowable, PageBreak, Image)
    from reportlab.lib.enums import TA_LEFT, TA_CENTER

    doc = SimpleDocTemplate(output_path, pagesize=A4,
                            leftMargin=2*cm, rightMargin=2*cm,
                            topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()
    def S(name, **kw):
        return ParagraphStyle(name, parent=styles["Normal"], **kw)

    S_TITLE = S("tt", fontSize=22, textColor=colors.HexColor("#000000"),
                spaceAfter=6, alignment=TA_CENTER, fontName="Helvetica-Bold")
    S_H1    = S("h1", fontSize=14, textColor=colors.HexColor("#000000"),
                spaceBefore=14, spaceAfter=4, fontName="Helvetica-Bold")
    S_H2    = S("h2", fontSize=11, textColor=colors.HexColor("#000000"),
                spaceBefore=8, spaceAfter=3, fontName="Helvetica-Bold")
    S_BODY  = S("bd", fontSize=9, textColor=colors.HexColor("#000000"),
                spaceAfter=4, leading=13)
    S_MONO  = S("mn", fontSize=8, textColor=colors.HexColor("#000000"),
                fontName="Courier", spaceAfter=2)
    S_MUTED = S("mu", fontSize=8, textColor=colors.HexColor("#000000"), spaceAfter=2)
    S_CTR   = S("ct", fontSize=9, textColor=colors.HexColor("#000000"), alignment=TA_CENTER)

    BG1 = colors.HexColor("#FFFFFF"); BG2 = colors.HexColor("#F6F6F6")
    BORDER = colors.HexColor("#B8B8B8"); CYAN_C = colors.HexColor("#000000")
    ORANGE_C = colors.HexColor("#000000")

    def tbl(data, col_widths, header_color=CYAN_C):
        t = Table(data, colWidths=col_widths)
        t.setStyle(TableStyle([
            ("BACKGROUND",  (0,0), (-1,0), BG2),
            ("TEXTCOLOR",   (0,0), (-1,0), header_color),
            ("FONTNAME",    (0,0), (-1,0), "Helvetica-Bold"),
            ("TEXTCOLOR",   (0,1), (-1,-1), colors.HexColor("#000000")),
            ("FONTSIZE",    (0,0), (-1,-1), 8),
            ("ROWBACKGROUNDS", (0,1), (-1,-1), [BG1, BG2]),
            ("BOX",         (0,0), (-1,-1), 0.5, BORDER),
            ("INNERGRID",   (0,0), (-1,-1), 0.25, BORDER),
            ("TOPPADDING",    (0,0), (-1,-1), 4),
            ("BOTTOMPADDING", (0,0), (-1,-1), 4),
        ]))
        return t

    story = []

    # Cover
    story += [Spacer(1, 3*cm), Paragraph(case.title or "Untitled Case", S_TITLE),
              Paragraph(case.case_number or "—", S_CTR), Spacer(1, 0.5*cm),
              HRFlowable(width="100%", color=BORDER), Spacer(1, 0.5*cm)]
    meta = _case_meta_rows(case)
    mt = tbl(meta, [4*cm, 13*cm])
    mt.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (0,-1), BG2),
        ("TEXTCOLOR",  (0,0), (0,-1), CYAN_C),
        ("FONTNAME",   (0,0), (0,-1), "Helvetica-Bold"),
        ("TEXTCOLOR",  (1,0), (1,-1), colors.HexColor("#000000")),
        ("FONTSIZE",   (0,0), (-1,-1), 9),
        ("ROWBACKGROUNDS", (0,0), (-1,-1), [BG1, BG2]),
        ("BOX",    (0,0), (-1,-1), 0.5, BORDER),
        ("INNERGRID", (0,0), (-1,-1), 0.25, BORDER),
        ("TOPPADDING",    (0,0), (-1,-1), 5),
        ("BOTTOMPADDING", (0,0), (-1,-1), 5),
    ]))
    story += [mt, PageBreak()]

    # Summary
    if case.description:
        story += [Paragraph("Summary", S_H1), HRFlowable(width="100%", color=BORDER),
                  Spacer(1, 0.2*cm),
                  Paragraph(case.description.replace("\n","<br/>"), S_BODY)]

    # Observables
    story += [Paragraph("ATT&amp;CK Observables", S_H1),
              HRFlowable(width="100%", color=BORDER), Spacer(1, 0.2*cm)]
    if case.observables:
        obs = [["T-Code","Technique Name","Confirmed","Added"]] + [
            [_norm_tcode(o.tcode), o.tname or "—", "Yes" if o.confirmed else "No", _ts(o.added_at)]
            for o in case.observables]
        story.append(tbl(obs, [3*cm, 9*cm, 2*cm, 4*cm]))
    else:
        story.append(Paragraph("No observables recorded.", S_MUTED))

    # Entities
    story += [Spacer(1, 0.4*cm), Paragraph("Entities Involved", S_H1),
              HRFlowable(width="100%", color=BORDER), Spacer(1, 0.2*cm)]
    if case.case_entities:
        ent = [["Type","Value","Role","Related Cases","Entity Description","Notes"]] + [
            [ce.entity.entity_type.value.replace("_"," ").title(),
             ce.entity.value, ce.role.value.title(), _entity_related_cases(case, ce), ce.entity.description or "—", ce.notes or "—"]
            for ce in case.case_entities]
        story.append(tbl(ent, [2.2*cm, 3.7*cm, 2.0*cm, 4.6*cm, 2.8*cm, 2.7*cm], header_color=ORANGE_C))
    else:
        story.append(Paragraph("No entities linked.", S_MUTED))

    # Attachments
    story += [Spacer(1, 0.4*cm), Paragraph("Attachments", S_H1),
              HRFlowable(width="100%", color=BORDER), Spacer(1, 0.2*cm)]
    if case.attachments:
        att_data = [["Filename", "Size", "Type", "Uploaded", "Description"]] + [
            [a.original_filename,
             f"{round(a.file_size/1024, 1)} KB" if a.file_size else "—",
             a.mime_type or "—",
             _ts(a.uploaded_at),
             a.description or "—"]
            for a in case.attachments
        ]
        story.append(tbl(att_data, [5*cm, 2*cm, 3*cm, 3.5*cm, 3.5*cm], header_color=ORANGE_C))
    else:
        story.append(Paragraph("No attachments.", S_MUTED))

    # BIAS
    story += [Spacer(1, 0.4*cm), Paragraph("BIAS Gap Analysis", S_H1),
              HRFlowable(width="100%", color=BORDER), Spacer(1, 0.2*cm)]
    if case.bias_result_json:
        try:
            result = json.loads(case.bias_result_json)
            arch = result.get("archetype", {})
            if arch.get("name"):
                story.append(Paragraph(
                    f"Archetype: <b>{arch['name']}</b>  (confidence: {arch.get('confidence',0):.0%})", S_BODY))
            graph_path = os.path.join(os.path.dirname(output_path), f"{case.case_number}_graph.png")
            graph_file = _render_bias_graph(case, graph_path)
            if graph_file and os.path.exists(graph_file):
                story += [Spacer(1, 0.2*cm), Image(graph_file, width=16.5*cm, height=8.8*cm), Spacer(1, 0.2*cm)]
            gap_num = 0
            for gap in result.get("gaps", []):
                conn_label = _connection_label(gap)
                if gap.get("is_gap", False):
                    gap_num += 1
                    title = f"{conn_label} {gap_num}: {gap.get('from_tcode','')} → {gap.get('to_tcode','')}"
                else:
                    title = f"{conn_label}: {gap.get('from_tcode','')} → {gap.get('to_tcode','')}"
                story.append(Paragraph(title, S_H2))
                if gap.get("missing_tactics"):
                    story.append(Paragraph(f"Missing tactics: {', '.join(gap.get('missing_tactics', []))}", S_MUTED))
                for c in gap.get("candidates", []):
                    tactics = ", ".join(c.get("tactics", [])) or "—"
                    platforms = ", ".join(c.get("platforms", [])) or "—"
                    story.append(Paragraph(
                        f"• {c.get('tcode','')}  {c.get('tname','')}  (score: {c.get('score',0):.2f})", S_MONO))
                    story.append(Paragraph(f"  Tactics: {tactics}  |  Platforms: {platforms}", S_MUTED))
                for ch in gap.get("chains", [])[:1]:
                    path = " → ".join(ch.get("path", []))
                    story.append(Paragraph(f"• Multi-hop: {path}  (score: {ch.get('score',0):.2f})", S_MONO))
        except Exception:
            story.append(Paragraph("BIAS result could not be parsed.", S_MUTED))
    else:
        story.append(Paragraph("BIAS analysis not yet run.", S_MUTED))

    # Checklist
    if case.checklist_json:
        try:
            items = json.loads(case.checklist_json)
            if items:
                story += [Spacer(1, 0.4*cm), Paragraph("Investigation Checklist", S_H1),
                          HRFlowable(width="100%", color=BORDER), Spacer(1, 0.2*cm)]
                for item in items:
                    done = item.get("done", False) if isinstance(item, dict) else False
                    text = item.get("text", str(item)) if isinstance(item, dict) else str(item)
                    col = "#3FB950" if done else "#8B949E"
                    mark = "✓" if done else "○"
                    story.append(Paragraph(f"{mark}  {text}", S_BODY))
        except Exception: pass

    # Notes
    if case.notes:
        story += [Spacer(1, 0.4*cm), Paragraph("Analyst Notes", S_H1),
                  HRFlowable(width="100%", color=BORDER), Spacer(1, 0.2*cm)]
        for note in case.notes:
            author = note.author.username if note.author else "unknown"
            story += [Paragraph(f"{_ts(note.created_at)}  —  {author}", S_MUTED),
                      Paragraph(note.body.replace("\n","<br/>"), S_BODY),
                      Spacer(1, 0.2*cm)]

    # Timeline
    if case.timeline:
        story += [Spacer(1, 0.4*cm), Paragraph("Case Timeline", S_H1),
                  HRFlowable(width="100%", color=BORDER), Spacer(1, 0.2*cm)]
        tl = [["Time","Event","Detail"]] + [
            [_ts(ev.occurred_at), ev.event_type.replace("_"," ").title(), ev.detail or "—"]
            for ev in sorted(case.timeline, key=lambda e: e.occurred_at)]
        story.append(tbl(tl, [4*cm, 4*cm, 9*cm]))

    # Footer
    story += [Spacer(1, 1*cm), HRFlowable(width="100%", color=BORDER),
              Paragraph(
                  f"Generated by LIST  ·  {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}",
                  S_CTR)]
    doc.build(story)
    return output_path


def build_docx(case: "Case", output_path: str) -> str:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(2)
        section.left_margin = section.right_margin = Cm(2.5)

    def heading(text, level=1, color=_LIGHT):
        p = doc.add_paragraph()
        r = p.add_run(text); r.bold = True
        r.font.size = Pt(16 - (level-1)*3)
        r.font.color.rgb = RGBColor(*color)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after  = Pt(4)
        return p

    def para(text, color=_LIGHT, size=10, mono=False):
        p = doc.add_paragraph()
        r = p.add_run(text); r.font.size = Pt(size)
        r.font.color.rgb = RGBColor(*color)
        if mono: r.font.name = "Courier New"
        p.paragraph_format.space_after = Pt(3)
        return p

    def table(headers, rows, col_widths=None):
        t = doc.add_table(rows=1+len(rows), cols=len(headers))
        t.style = "Table Grid"
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]; c.text = h
            r = c.paragraphs[0].runs[0]; r.bold = True
            r.font.color.rgb = RGBColor(*_LIGHT); r.font.size = Pt(9)
        for ri, row in enumerate(rows):
            for ci, val in enumerate(row):
                c = t.rows[ri+1].cells[ci]; c.text = str(val)
                r = c.paragraphs[0].runs[0]
                r.font.size = Pt(9); r.font.color.rgb = RGBColor(*_LIGHT)
        if col_widths:
            for row in t.rows:
                for ci, w in enumerate(col_widths):
                    row.cells[ci].width = Cm(w)
        return t

    # Cover
    for txt, sz in [(case.title or "Untitled Case", 24), (case.case_number or "—", 12)]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(txt); r.bold = True
        r.font.size = Pt(sz); r.font.color.rgb = RGBColor(*_LIGHT)
    doc.add_paragraph()
    table(["Field","Value"], _case_meta_rows(case), col_widths=[4, 13])
    doc.add_page_break()

    if case.description:
        heading("Summary"); para(case.description)

    heading("ATT&CK Observables")
    if case.observables:
        table(["T-Code","Technique Name","Confirmed","Added"],
              [[_norm_tcode(o.tcode), o.tname or "—", "Yes" if o.confirmed else "No", _ts(o.added_at)] for o in case.observables],
              col_widths=[3, 9, 2, 4])
    else: para("No observables recorded.", color=_MUTED)

    heading("Entities Involved", color=_LIGHT)
    if case.case_entities:
        table(["Type","Value","Role","Related Cases","Entity Description","Notes"],
              [[ce.entity.entity_type.value.replace("_"," ").title(),
                ce.entity.value, ce.role.value.title(), _entity_related_cases(case, ce), ce.entity.description or "—", ce.notes or "—"]
               for ce in case.case_entities],
              col_widths=[2.2, 3.7, 2.0, 4.6, 2.8, 2.7])
    else: para("No entities linked.", color=_MUTED)

    heading("Attachments", level=1, color=_LIGHT)
    if case.attachments:
        table(["Filename", "Size", "Type", "Uploaded", "Description"],
              [[a.original_filename,
                f"{round(a.file_size/1024, 1)} KB" if a.file_size else "—",
                a.mime_type or "—",
                _ts(a.uploaded_at),
                a.description or "—"]
               for a in case.attachments],
              col_widths=[5, 2, 3, 3.5, 3.5])
    else:
        para("No attachments.", color=_MUTED)

    heading("BIAS Gap Analysis")
    if case.bias_result_json:
        try:
            result = json.loads(case.bias_result_json)
            arch = result.get("archetype", {})
            if arch.get("name"):
                para(f"Archetype: {arch['name']}  (confidence: {arch.get('confidence',0):.0%})", color=_LIGHT)
            graph_path = os.path.join(os.path.dirname(output_path), f"{case.case_number}_graph.png")
            graph_file = _render_bias_graph(case, graph_path)
            if graph_file and os.path.exists(graph_file):
                doc.add_picture(graph_file, width=Cm(16.5))
            gap_num = 0
            for gap in result.get("gaps", []):
                conn_label = _connection_label(gap)
                if gap.get("is_gap", False):
                    gap_num += 1
                    title = f"{conn_label} {gap_num}: {gap.get('from_tcode','')} → {gap.get('to_tcode','')}"
                else:
                    title = f"{conn_label}: {gap.get('from_tcode','')} → {gap.get('to_tcode','')}"
                heading(title, level=2, color=_LIGHT)
                if gap.get("missing_tactics"):
                    para(f"Missing tactics: {', '.join(gap.get('missing_tactics', []))}", color=_MUTED, size=8)
                for c in gap.get("candidates", []):
                    para(f"  • {c.get('tcode','')}  {c.get('tname','')}  score: {c.get('score',0):.2f}", color=_LIGHT, mono=True)
                    para(f"    Tactics: {', '.join(c.get('tactics', [])) or '—'}  |  Platforms: {', '.join(c.get('platforms', [])) or '—'}",
                         color=_MUTED, size=8)
                for ch in gap.get("chains", [])[:1]:
                    para(f"  • Multi-hop: {' → '.join(ch.get('path',[]))}  score: {ch.get('score',0):.2f}", color=_LIGHT, mono=True)
        except Exception: para("BIAS result could not be parsed.", color=_MUTED)
    else: para("BIAS analysis not yet run.", color=_MUTED)

    if case.checklist_json:
        try:
            items = json.loads(case.checklist_json)
            if items:
                heading("Investigation Checklist", level=2)
                for item in items:
                    done = item.get("done", False) if isinstance(item, dict) else False
                    text = item.get("text", str(item)) if isinstance(item, dict) else str(item)
                    para(f"{'✓' if done else '○'}  {text}", color=_LIGHT)
        except Exception: pass

    if case.notes:
        heading("Analyst Notes")
        for note in case.notes:
            author = note.author.username if note.author else "unknown"
            para(f"{_ts(note.created_at)}  —  {author}", color=_MUTED, size=8)
            para(note.body)
            doc.add_paragraph()

    if case.timeline:
        heading("Case Timeline")
        table(["Time","Event","Detail"],
              [[_ts(ev.occurred_at), ev.event_type.replace("_"," ").title(), ev.detail or "—"]
               for ev in sorted(case.timeline, key=lambda e: e.occurred_at)],
              col_widths=[4, 4, 9])

    doc.add_paragraph()
    fp = doc.add_paragraph(); fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run(f"Generated by LIST  ·  {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}")
    fr.font.size = Pt(8); fr.font.color.rgb = RGBColor(*_MUTED)
    doc.save(output_path)
    return output_path
